from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
from docx.enum.section import WD_ORIENT

doc = Document()

style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(6)

section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width = Inches(11)
section.page_height = Inches(8.5)

table = doc.add_table(rows=2, cols=9)
table.style = 'Table Grid'

headers = ["Код ЗВ", "Наименование загрязняющего вещества", "ПДК максимальная разовая мг/м3", "ПДК среднесуточная мг/м3", "ПДК среднегодовая мг/м3", "ОБУВ мг/м3",
           "Класс опасности", "Выброс вещества, г/с", "Суммарный выброс вещества, т/год"]

for i, header in enumerate(headers):
    table.cell(0, i).text = header

doc.save(r"calculations\tables\razdel2_table.docx")