from docx import Document

def fill_3_8_table(number, izav_type, izav_amount, izav_speed, fuel_type, work_time_season, work_time_year, nitric_ocide_1_year_warm, nitric_ocide_year_warm, sulfur_year_warm, CO_year_warm, petrol_year_warm, nitric_ocide_1_sec_warm, nitric_ocide_sec_warm, sulfur_sec_warm, CO_sec_warm, petrol_sec_warm):
    doc = Document("calculations/tables/3_8.docx")
    table = doc.tables[0]
    table.cell(2,0).text = str(number)
    table.cell(2,1).text = izav_type
    table.cell(2,2).text = str(izav_amount)
    table.cell(2,3).text = str(izav_speed)
    table.cell(2,4).text = fuel_type
    table.cell(2,5).text = str(work_time_season)
    table.cell(2,6).text = str(work_time_year)

    table.cell(2,7).text = "(0301) Азота диоксид\n (0304) Азот (II) оксид\n 0330 Сера диоксид\n 0337 Углерода оксид\n 2704 Бензин (нефтяной, малосернистый) /в пересчете на углерод/\n"
    table.cell(2,8).text = f"{nitric_ocide_1_sec_warm}\n{nitric_ocide_sec_warm}\n {sulfur_sec_warm}\n {CO_sec_warm}\n {petrol_sec_warm}\n"
    table.cell(2,9).text = f"{nitric_ocide_1_year_warm}\n{nitric_ocide_year_warm}\n {sulfur_year_warm}\n {CO_year_warm}\n {petrol_year_warm}"
    table.cell(2,10).text = "0001"
    table.cell(3,0).text = "Всего"
    table.cell(3,2).text = "1"
    table.cell(3,9).text = str(nitric_ocide_1_year_warm + nitric_ocide_year_warm + sulfur_year_warm + CO_year_warm + petrol_year_warm)
    table.add_row()
    table.cell(4,0).merge(table.cell(4,10))
    table.cell(4,0).text = """Примечание. Список использованных расчетных методик: 0001 - Методика проведения инвентаризации выбросов ЗВ в атмосферу для автотранспортных предприя-тий". - М., НИИАТ, 1998г. с учётом "Дополнений к методике" ( М., НИИАТ, 1999г.) и Методики…для баз дорожной техники". - М., НИИАТ, 1999г """
    doc.save("calculations/tables/razdel3/3_8.docx")    